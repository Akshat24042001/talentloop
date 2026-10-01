"""Selected candidates as an .xlsx for HROne's employee import.

TalentLoop does not ship HROne's column names: each company copies the headers from its own HROne import template
into Settings (hrone_columns: [{header, field}]) and maps every header to one of the TalentLoop fields below.
Headers mapped to "" are written empty, so the sheet still matches the template column for column.
"""
import io
import time
from datetime import datetime

from . import db, jd_schema, refs

FIELDS = {
    "full_name": "Full name", "first_name": "First name", "last_name": "Last name", "email": "Email", "phone": "Phone",
    "location": "Current city", "current_company": "Current company", "current_title": "Current job title", "years": "Total experience (years)",
    "notice_days": "Notice period (days)", "expected_salary": "Expected salary", "college": "College", "degree": "Degree",
    "graduation_year": "Year of passing", "linkedin": "LinkedIn", "job_title": "Job title (the role hired for)", "department": "Department",
    "job_location": "Job location", "employment_type": "Employment type", "job_ref": "Job reference", "candidate_ref": "Candidate reference",
    "selected_on": "Selected on (date)", "stage": "Stage", "source": "Source", "rating": "Team rating",
}
DEFAULT_COLUMNS = [{"header": FIELDS[k], "field": k} for k in ("full_name", "email", "phone", "job_title", "department", "job_location", "selected_on")]


def _split(name: str) -> tuple[str, str]:
    parts = (name or "").split()
    return (parts[0] if parts else "", " ".join(parts[1:]))


def value(field: str, a: db.Application, c: db.Candidate, j: db.Job) -> object:
    p, f = c.profile or {}, j.fields or {}
    edu = (p.get("education") or [{}])[0] if isinstance(p.get("education"), list) and p.get("education") else {}
    v = {
        "full_name": c.name, "first_name": _split(c.name)[0], "last_name": _split(c.name)[1], "email": c.email, "phone": c.phone,
        "location": c.location, "current_company": c.current_company, "current_title": c.headline, "years": c.years,
        "notice_days": int(c.notice_days) if c.notice_days is not None else None, "expected_salary": c.expected_salary, "college": c.college,
        "degree": p.get("degree") or edu.get("degree"), "graduation_year": p.get("graduation_year") or edu.get("year"), "linkedin": p.get("linkedin"),
        "job_title": j.title, "department": j.department, "job_location": jd_schema.location_text(f),
        "employment_type": f.get("employment_type"), "job_ref": refs.job_ref(j), "candidate_ref": refs.cand_ref(c),
        "selected_on": datetime.fromtimestamp(a.decided_at or a.updated_at).strftime("%d-%m-%Y") if (a.decided_at or a.updated_at) else None,
        "stage": a.stage, "source": a.source, "rating": a.rating,
    }.get(field)
    return "" if v is None else v


def workbook(columns: list[dict], rows: list[tuple[db.Application, db.Candidate, db.Job]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    cols = [c for c in (columns or DEFAULT_COLUMNS) if c.get("header")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append([c["header"] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for a, c, j in rows:
        ws.append([value(col.get("field") or "", a, c, j) for col in cols])
        for cell in ws[ws.max_row]:            # candidate-typed text is data, never a formula
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(col["header"]) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(job: db.Job | None) -> str:
    return f"hrone-import-{(refs.job_ref(job) if job else 'selected')}-{time.strftime('%Y%m%d')}.xlsx"
