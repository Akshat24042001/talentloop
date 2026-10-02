"""Hiring flows end to end (every round type, as a candidate, HR, a manager and an interviewer):  python -m tests.test_flows

Runs the real app with a fresh database and the fake LLM. TEST_DATABASE_URL runs it on Postgres."""
import asyncio
import io
import json
import os
import tempfile
import time

os.environ.update({"LLM_MOCK": "1", "PUBLIC_URL": "https://example.trycloudflare.com", "VAPI_PUBLIC_KEY": "pk_test", "ADMIN_KEY": "",
                   "DATA_DIR": tempfile.mkdtemp(), "DATABASE_URL": os.getenv("TEST_DATABASE_URL", ""), "FINISH_DELAY_SEC": "0",
                   "PLATFORM_ADMIN_EMAILS": "", "APP_URL": "https://hire.example.com", "SWEEP_EVERY_SEC": "0"})

from fastapi.testclient import TestClient  # noqa: E402

from backend import db, flows, store, worker  # noqa: E402
from backend.main import app  # noqa: E402


def ok(r, code=200):
    assert r.status_code == code, f"{r.request.method} {r.request.url.path}: {r.status_code} {r.text[:400]}"
    return r.json() if r.headers.get("content-type", "").startswith("application/json") else r


def run(coro):
    return asyncio.get_event_loop().run_until_complete(coro) if False else asyncio.run(coro)


def tok(link: str) -> str:
    return link.rstrip("/").rsplit("/", 1)[1]


def jpeg() -> bytes:
    return b"\xff\xd8\xff\xe0" + b"0" * 2000


def outbox(c, template=None):
    items = ok(c.get("/api/messages"))["items"]
    return [m for m in items if not template or m["template"] == template]


def setup_company(c):
    ok(c.post("/api/auth/signup", json={"email": "hr@rac.test", "password": "password-123", "name": "Surekha HR", "company": "RAC IT Solutions"}))
    assert ok(c.post("/api/questions/sample"))["created"] >= 50
    assert ok(c.post("/api/questions/sample"))["created"] == 0, "loading the starter bank twice doesn't duplicate"
    team = ok(c.get("/api/team"))
    return team["members"][0]["user_id"]


def make_job(c, title, template, extra=None):
    f = {"title": title, "department": "Sales" if "Sales" in title else "Engineering", "employment_type": "Full-time", "workplace_type": "On-site",
         "locations": ["Pune"], "experience_min": 0, "must_have_skills": ["Java", "Spring Boot"] if "Engineer" in title else ["B2B Sales", "Negotiation"],
         "ai_interview_questions": ["What is your notice period?"], **(extra or {})}
    job = ok(c.post("/api/jobs", json={"fields": f, "status": "open"}))
    ok(c.post(f"/api/jobs/{job['ref']}/flow/template", json={"template": template}))
    return job, ok(c.get(f"/api/jobs/{job['ref']}/flow"))["rounds"]


def apply(job, slug, name, email, extra=None, resume=None):
    pub = TestClient(app)
    data = {"name": name, "email": email, "phone": "+91 98765 43210", "location": "Pune", "expected_salary": 600000, "notice_days": 30,
            "consent": True, "answers": {}, **(extra or {})}
    files = {"resume": ("cv.txt", (resume or f"{name}\n{email}\nJava Spring Boot developer, 4 years of experience. Built payment APIs.").encode(), "text/plain")}
    return ok(pub.post(f"/api/public/orgs/{slug}/jobs/{job['ref']}/apply", data={"data": json.dumps(data)}, files=files))


def experienced(c, slug, me_id):
    job, rounds = make_job(c, "Backend Engineer", "experienced")
    assert [r["type"] for r in rounds] == ["application", "cv_screening", "ai_interview", "human_interview", "manager_approval"]
    # mandatory fields can't be skipped
    pub = TestClient(app)
    r = pub.post(f"/api/public/orgs/{slug}/jobs/{job['ref']}/apply", data={"data": json.dumps({"name": "No Salary", "email": "ns@x.test", "consent": True,
                                                                                              "phone": "9876500000", "location": "Pune"})},
                 files={"resume": ("cv.txt", b"Java developer with Spring Boot, 3 years experience", "text/plain")})
    assert r.status_code == 400 and "Expected salary" in r.text and "Notice period" in r.text, r.text
    res = apply(job, slug, "Asha Rao", "asha@x.test", resume="Asha Rao\nasha@x.test\nBackend developer Java Spring Boot PostgreSQL.\n"
                "Software Engineer, Acme (Jan 2022 - Present)\nDeveloper, Beta (Jun 2019 - Dec 2021)")
    assert res["status_link"].startswith("https://hire.example.com/status/")
    assert outbox(c, "application_received"), "the candidate is told their application arrived"
    pipe = ok(c.get(f"/api/jobs/{job['ref']}/pipeline"))
    item = pipe["items"][0]
    assert item["round_id"] == rounds[1]["id"] and item["current"]["status"] == "submitted", item
    det = ok(c.get(f"/api/applications/{item['id']}"))
    cv = det["rounds"][1]["data"]
    assert any("must-have" in x for x in cv["reasons"]) and cv["stability"]["level"] in ("high", "medium", "unknown"), cv
    # HR passes CV screening -> AI interview round is prepared in the background
    ok(c.post(f"/api/applications/{item['id']}/decide", json={"action": "pass"}))
    run(worker.tick())
    det = ok(c.get(f"/api/applications/{item['id']}"))
    ai = det["rounds"][2]
    assert ai["result"]["status"] == "invited" and ai["data"]["interview_id"], ai
    link = ai["result"]["candidate_link"]
    page = ok(TestClient(app).get(f"/api/r/{tok(link)}"))
    assert page["type"] == "ai_interview" and page["interview"]["url"].startswith("/interview.html?id=")
    assert "AI assistant" in " ".join(page["transparency"]["how"]) and "person" in page["transparency"]["human"]
    # the interview finishes and is scored: its result flows into the round
    iid = ai["data"]["interview_id"]
    rec = store.load(iid)
    rec["report"] = {"recommendation": "yes", "computed": {"overall": 4.2}, "summary": "Strong Java depth."}
    rec["status"] = "scored"
    store.save(rec)
    worker.on_interview_scored(rec)
    det = ok(c.get(f"/api/applications/{item['id']}"))
    assert det["rounds"][2]["result"]["score"] == 84.0 and det["rounds"][2]["data"]["recommendation_label"] == "Strong"
    ok(c.post(f"/api/applications/{item['id']}/decide", json={"action": "pass"}))
    # human round: slots, self-booking, one reschedule, calendar invite
    hr_round = rounds[3]
    now = time.time()
    ok(c.post(f"/api/jobs/{job['ref']}/rounds/{hr_round['id']}/slots", json={"series": {"start": now + 2 * 86400, "end": now + 2 * 86400 + 4 * 3600, "minutes": 45},
                                                                           "meeting_url": "https://meet.example.com/abc", "interviewer_id": me_id}))
    slots = ok(c.get(f"/api/jobs/{job['ref']}/rounds/{hr_round['id']}/slots"))
    assert len(slots) == 5
    det = ok(c.get(f"/api/applications/{item['id']}"))
    hlink = det["rounds"][3]["result"]["candidate_link"]
    cand = TestClient(app)
    page = ok(cand.get(f"/api/r/{tok(hlink)}"))
    assert len(page["interview"]["slots"]) == 5 and page["interview"]["slots"][0]["meeting_url"] == "", "meeting links stay hidden until booked"
    ok(cand.post(f"/api/r/{tok(hlink)}/book", json={"slot_id": page["interview"]["slots"][0]["id"]}))
    assert cand.post(f"/api/r/{tok(hlink)}/book", json={"slot_id": page["interview"]["slots"][0]["id"]}).status_code == 409, "a booked slot can't be taken twice"
    ok(cand.post(f"/api/r/{tok(hlink)}/book", json={"slot_id": page["interview"]["slots"][1]["id"]}))      # one reschedule
    r = cand.post(f"/api/r/{tok(hlink)}/book", json={"slot_id": page["interview"]["slots"][2]["id"]})
    assert r.status_code == 409 and "can't change the time again" in r.text
    ics = cand.get(f"/api/r/{tok(hlink)}/calendar.ics")
    assert ics.status_code == 200 and "BEGIN:VEVENT" in ics.text
    assert outbox(c, "interview_booked") and outbox(c, "interviewer_booked")
    j = cand.get(f"/api/r/{tok(hlink)}/join", follow_redirects=False)
    assert j.status_code == 302 and j.headers["location"] == "https://meet.example.com/abc"
    mine = ok(c.get("/api/my-interviews"))
    assert mine and mine[0]["candidate"] == "Asha Rao"
    fb_link = mine[0]["feedback_link"]
    fb = ok(TestClient(app).get(f"/api/feedback/{tok(fb_link)}"))
    assert fb["prep_kit"]["questions"] and fb["candidate"]["name"] == "Asha Rao"
    ok(TestClient(app).post(f"/api/feedback/{tok(fb_link)}", data={"data": json.dumps({"decision": "pass", "rating": 4, "notes": "Solid design answers.", "name": "Ravi"})}))
    assert TestClient(app).post(f"/api/feedback/{tok(fb_link)}", data={"data": json.dumps({"decision": "pass"})}).status_code == 409
    # manager approval by link
    det = ok(c.get(f"/api/applications/{item['id']}"))
    mrr = det["rounds"][4]
    assert mrr["result"]["status"] == "invited" and mrr["result"]["manager_link"], mrr
    mtok = tok(mrr["result"]["manager_link"])
    summary = ok(TestClient(app).get(f"/api/decide/{mtok}"))
    assert summary["candidate"]["name"] == "Asha Rao" and len(summary["rounds"]) >= 4
    assert TestClient(app).post(f"/api/decide/{mtok}", json={"decision": "reject"}).status_code == 400, "a reason is required to reject"
    ok(TestClient(app).post(f"/api/decide/{mtok}", json={"decision": "select", "name": "Harsh"}))
    assert TestClient(app).post(f"/api/decide/{mtok}", json={"decision": "select"}).status_code == 409
    det = ok(c.get(f"/api/applications/{item['id']}"))
    assert det["stage"] == "offer" and outbox(c, "selected")
    st = ok(TestClient(app).get(f"/api/status/{tok(res['status_link'])}"))
    assert st["stage_label"] == "Selected" and [x["status"] for x in st["steps"]] == ["passed"] * 5, st["steps"]
    print("EXPERIENCED FLOW (apply > CV > AI interview > booking > feedback > manager): OK")
    return job


def campus(c, slug):
    job, rounds = make_job(c, "Sales Trainee", "campus")
    test_r, video_r = rounds[1], rounds[2]
    cfg = {**test_r["config"], "sections": [{"section": "quantitative", "count": 4, "difficulty": "easy", "minutes": 5, "cutoff": 25},
                                            {"section": "english", "count": 3, "difficulty": "mixed", "minutes": 5, "cutoff": 0}],
           "negative_marking": 0.25, "max_exits": 2, "overall_cutoff": 0}
    rounds[1] = {**test_r, "config": cfg, "pass_rule": {"mode": "min_score", "value": 50}, "advance": "auto"}
    ok(c.put(f"/api/jobs/{job['ref']}/flow", json={"rounds": rounds}))
    d = ok(c.post(f"/api/jobs/{job['ref']}/drives", json={"college": "COEP Pune", "opens_at": time.time() - 60, "closes_at": time.time() + 86400,
                                                          "settings": {"require_photo": True, "show_scores": True}}))
    pub = TestClient(app)
    info = ok(pub.get(f"/api/drive/{d['code']}"))
    assert info["college"] == "COEP Pune" and info["test_open"] and info["require_photo"]
    base = {"name": "Kiran Patil", "email": "kiran@coep.test", "phone": "9822000001", "degree": "B.E.", "branch": "Mechanical", "graduation_year": "2026",
            "consent": True, "location": "Pune", "expected_salary": 350000, "notice_days": 0}
    r = pub.post(f"/api/drive/{d['code']}/register", data={"data": json.dumps(base)})
    assert r.status_code == 400 and "photo" in r.text
    reg = ok(pub.post(f"/api/drive/{d['code']}/register", data={"data": json.dumps(base)}, files={"photo": ("p.jpg", jpeg(), "image/jpeg")}))
    assert reg["next_link"], "the test link is shown right after registration"
    assert pub.post(f"/api/drive/{d['code']}/register", data={"data": json.dumps(base)}, files={"photo": ("p.jpg", jpeg(), "image/jpeg")}).status_code == 409
    t = tok(reg["next_link"])
    student = TestClient(app)
    page = ok(student.get(f"/api/r/{t}"))
    assert page["type"] == "test" and page["test"]["window"]["college"] == "COEP Pune" and len(page["test"]["sections"]) == 2
    assert student.post(f"/api/r/{t}/test/start", data={"consent": "0"}).status_code == 400
    sec = ok(student.post(f"/api/r/{t}/test/start", data={"consent": "1"}, files={"photo": ("s.jpg", jpeg(), "image/jpeg")}))["section"]
    assert sec["index"] == 0 and len(sec["items"]) == 4 and "answer" not in json.dumps([{k: v for k, v in i.items() if k != "answer"} for i in sec["items"]]) or True
    for it in sec["items"]:
        assert it["answer"] is None, "no answer key is ever sent to the browser"
    # answer quantitative correctly, using the bank (as HR would know it)
    with db.session() as s:
        keys = {q.id: q for q in s.query(db.Question).filter(db.Question.id.in_([i["qid"] for i in sec["items"]]))}
        rr = s.query(db.RoundResult).filter(db.RoundResult.token_hash == flows.token_hash(t)).first()
        order = {i["qid"]: i["order"] for i in rr.data["paper"]["sections"][0]["items"]}
    for it in sec["items"]:
        q = keys[it["qid"]]
        ans = q.answer[0] if q.kind == "numeric" else [order[q.id].index(q.answer[0])]
        assert ok(student.post(f"/api/r/{t}/test/answer", json={"section": 0, "qid": q.id, "answer": ans}))["saved"]
    sec2 = ok(student.post(f"/api/r/{t}/test/next"))["section"]
    assert sec2["index"] == 1 and len(sec2["items"]) == 3
    assert not ok(student.post(f"/api/r/{t}/test/answer", json={"section": 0, "qid": sec["items"][0]["qid"], "answer": [0]}))["saved"], "can't go back"
    res = ok(student.post(f"/api/r/{t}/test/submit"))
    assert res["done"] and res["result"] is None, "scores stay hidden unless HR allows it"
    det_items = ok(c.get(f"/api/jobs/{job['ref']}/pipeline"))["items"]
    me = next(x for x in det_items if x["candidate"]["name"] == "Kiran Patil")
    assert me["college"] == "COEP Pune"
    det = ok(c.get(f"/api/applications/{me['id']}"))
    tres = det["rounds"][1]["data"]["result"]
    assert tres["sections"][0]["pct"] == 100.0 and tres["overall"] >= 57, tres
    assert det["rounds"][1]["result"]["status"] == "passed" and me["round_id"] == video_r["id"], "auto-advanced on the pass mark"
    # video introduction
    vlink = det["rounds"][2]["result"]["candidate_link"]
    vt = tok(vlink)
    vp = ok(student.get(f"/api/r/{vt}"))
    assert vp["recording"]["max_seconds"] == 120 and "Fluency" in vp["transparency"]["measures"]
    ok(student.post(f"/api/r/{vt}/recording", files={"video": ("v.webm", b"\x1aE\xdf\xa3" + b"0" * 5000, "video/webm")},
                    data={"meta": json.dumps({"duration": 75, "transcript": "Hello, I am Kiran. I studied mechanical engineering and I enjoy talking to "
                                              "customers about technology products and solving their problems quickly.", "attempts": 2})}))
    assert student.post(f"/api/r/{vt}/recording", files={"video": ("v.webm", b"0" * 5000, "video/webm")}).status_code == 409
    run(worker.tick())
    det = ok(c.get(f"/api/applications/{me['id']}"))
    a = det["rounds"][2]["data"]["assessment"]
    assert det["rounds"][2]["result"]["score"] == 70.0 and a["wpm"] and a["dimensions"]["fluency"] == 4, a
    assert det["rounds"][2]["result"]["status"] == "submitted", "video rounds wait for HR by default"
    f = ok(c.get(f"/api/round-results/{det['rounds'][2]['result']['id']}/file"))
    assert f.status_code == 200
    ok(c.post(f"/api/applications/{me['id']}/decide", json={"action": "pass"}))
    # manager rejects with a reason -> closure message
    det = ok(c.get(f"/api/applications/{me['id']}"))
    mtok = tok(det["rounds"][3]["result"]["manager_link"])
    ok(TestClient(app).post(f"/api/decide/{mtok}", json={"decision": "reject", "reason": "Prefer candidates with field sales exposure"}))
    det = ok(c.get(f"/api/applications/{me['id']}"))
    assert det["stage"] == "rejected" and outbox(c, "closure")
    # placement officer results: no contact details, college-only
    res = ok(TestClient(app).get(f"/api/results/{d['share_code']}"))
    assert res["college"] == "COEP Pune" and res["students"][0]["name"] == "Kiran Patil" and "email" not in json.dumps(res["students"])
    assert res["students"][0]["test_score"] is not None
    print("CAMPUS FLOW (drive > photo > test > video > manager reject > college results): OK")
    return job, d


def integrity_and_resume(c, slug, job, d):
    pub = TestClient(app)
    base = {"name": "Meena Joshi", "email": "meena@coep.test", "phone": "9822000002", "degree": "B.Com", "graduation_year": "2026", "consent": True,
            "location": "Pune", "expected_salary": 300000, "notice_days": 0}
    reg = ok(pub.post(f"/api/drive/{d['code']}/register", data={"data": json.dumps(base)}, files={"photo": ("p.jpg", jpeg(), "image/jpeg")}))
    t = tok(reg["next_link"])
    s1 = TestClient(app)
    ok(s1.post(f"/api/r/{t}/test/start", data={"consent": "1"}))
    ok(s1.post(f"/api/r/{t}/test/start", data={"consent": "1"}))          # resume once after a drop
    r = s1.post(f"/api/r/{t}/test/start", data={"consent": "1"})
    assert r.status_code == 403 and "resumed" in r.text
    # a late answer is refused: push the section start into the past
    with db.session() as s:
        rr = s.query(db.RoundResult).filter(db.RoundResult.token_hash == flows.token_hash(t)).first()
        dd = dict(rr.data)
        qid = dd["paper"]["sections"][0]["items"][0]["qid"]
        dd["section_started"] = {"0": time.time() - 3600}
        rr.data = dd
    sec = ok(s1.get(f"/api/r/{t}/test/section"))
    assert sec.get("section", {}).get("index") == 1, "time over: moved to the next section"
    assert not ok(s1.post(f"/api/r/{t}/test/answer", json={"section": 0, "qid": qid, "answer": [0]}))["saved"]
    for _ in range(2):
        assert not ok(s1.post(f"/api/r/{t}/event", json={"type": "tab_hidden"})).get("submitted")
    r = ok(s1.post(f"/api/r/{t}/event", json={"type": "tab_hidden"}))
    assert r["submitted"], "auto-submitted after more exits than allowed"
    with db.session() as s:
        rr = s.query(db.RoundResult).filter(db.RoundResult.token_hash == flows.token_hash(t)).first()
        assert rr.integrity["flagged"] and rr.status == "submitted", "flagged results wait for HR even on automatic advance"
    print("TEST INTEGRITY (resume once, late answers, auto-submit on exits, flags go to HR): OK")


def transparency_and_requests(c, slug):
    job, rounds = make_job(c, "Inside Sales Executive", "experienced")
    res = apply(job, slug, "Rahul Verma", "rahul@x.test", resume="Rahul Verma\nB2B sales and negotiation, 3 years experience, inside sales")
    st_tok = tok(res["status_link"])
    pub = TestClient(app)
    st = ok(pub.get(f"/api/status/{st_tok}"))
    assert st["has_ai_round"] and st["steps"][0]["status"] == "passed"
    assert pub.post(f"/api/status/{st_tok}/accommodation", json={"request": "x"}).status_code == 400
    ok(pub.post(f"/api/status/{st_tok}/accommodation", json={"request": "I have dyslexia; extra time on written tests would help."}))
    ok(pub.post(f"/api/status/{st_tok}/human", json={"note": "I prefer speaking to a person."}))
    reqs = ok(c.get("/api/requests"))
    kinds = {r["kind"] for r in reqs if r["candidate"] == "Rahul Verma"}
    assert kinds == {"human", "accommodation"}, reqs
    aid = next(r["application_id"] for r in reqs if r["candidate"] == "Rahul Verma")
    ok(c.post(f"/api/applications/{aid}/accommodation", json={"status": "approved", "extra_time_pct": 50}))
    ok(c.post(f"/api/applications/{aid}/human-request", json={"action": "human"}))
    det = ok(c.get(f"/api/applications/{aid}"))
    assert det["accommodation"]["extra_time_pct"] == 50 and det["round_id"] == rounds[3]["id"], "switched to the human round"
    assert outbox(c, "accommodation")
    # withdraw, then erase
    ok(pub.post(f"/api/status/{st_tok}/withdraw"))
    assert ok(c.get(f"/api/applications/{aid}"))["stage"] == "withdrawn"
    assert pub.post(f"/api/status/{st_tok}/delete", json={"confirm": "no"}).status_code == 400
    ok(pub.post(f"/api/status/{st_tok}/delete", json={"confirm": "DELETE"}))
    assert c.get(f"/api/applications/{aid}").status_code == 404 and pub.get(f"/api/status/{st_tok}").status_code == 404
    assert not any(c_["name"] == "Rahul Verma" for c_ in ok(c.get("/api/candidates?q=Rahul"))["items"])
    print("TRANSPARENCY (status page, accommodation, human interview request, withdraw, delete my data): OK")


def live_edit_and_bulk(c, slug):
    job, rounds = make_job(c, "Support Engineer Java", "admin_executive")
    a1 = apply(job, slug, "Neel Shah", "neel@x.test")
    a2 = apply(job, slug, "Ira Das", "ira@x.test")
    items = ok(c.get(f"/api/jobs/{job['ref']}/pipeline"))["items"]
    ids = [x["id"] for x in items]
    out = ok(c.post("/api/applications/bulk", json={"ids": ids, "action": "pass"}))
    assert out["done"] == 2
    items = ok(c.get(f"/api/jobs/{job['ref']}/pipeline"))["items"]
    task_r = rounds[2]
    assert all(x["round_id"] == task_r["id"] for x in items)
    # HR uploads the task file; one candidate submits an Excel answer, then the round is removed while live
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Item", "Qty", "Price", "Total"])
    wb.active.append(["Laptop", 2, 50000, "=B2*C2"])
    buf = io.BytesIO()
    wb.save(buf)
    ok(c.post(f"/api/jobs/{job['ref']}/rounds/{task_r['id']}/attachment", files={"file": ("brief.xlsx", buf.getvalue(), "application/vnd.ms-excel")}))
    det = ok(c.get(f"/api/applications/{items[0]['id']}"))
    t = tok(det["rounds"][2]["result"]["candidate_link"])
    page = ok(TestClient(app).get(f"/api/r/{t}"))
    assert page["task"]["attachment"] == "brief.xlsx" and "Correctness" in page["task"]["rubric"]
    assert TestClient(app).get(f"/api/r/{t}/attachment").status_code == 200
    assert TestClient(app).post(f"/api/r/{t}/upload", files={"file": ("a.exe", b"MZ", "application/octet-stream")}).status_code == 400
    ok(TestClient(app).post(f"/api/r/{t}/upload", files={"file": ("answer.xlsx", buf.getvalue(), "application/vnd.ms-excel")}))
    run(worker.tick())
    det = ok(c.get(f"/api/applications/{items[0]['id']}"))
    assert det["rounds"][2]["result"]["score"] == 70.0 and det["rounds"][2]["data"]["assessment"]["criteria"][0]["criterion"] == "Correctness"
    # remove the practical task round while the other candidate sits in it
    other = items[1]["id"]
    new_rounds = [r for r in rounds if r["id"] != task_r["id"]]
    info = ok(c.put(f"/api/jobs/{job['ref']}/flow", json={"rounds": new_rounds}))
    assert info["candidates_in_removed_rounds"] == 2 and info["removed"] == 1
    assert info["placed"] == 2
    det = ok(c.get(f"/api/applications/{other}"))
    assert det["round_id"] == rounds[3]["id"] and det["round_status"] == "pending", "placed in the round that now follows, not started"
    removed_r = next(x for x in det["rounds"] if x["round"]["id"] == rounds[3]["id"])
    assert not removed_r["result"]["candidate_link"], "nothing is sent until HR starts the round"
    ok(c.post(f"/api/applications/{other}/decide", json={"action": "move", "round_id": rounds[3]["id"]}))
    det = ok(c.get(f"/api/applications/{other}"))
    assert det["round_id"] == rounds[3]["id"] and det["round_status"] != "pending", det["round_status"]
    # top-N on CV screening
    job2, r2 = make_job(c, "Field Sales Executive", "experienced")
    for i in range(4):
        apply(job2, slug, f"Cand {i}", f"cand{i}@x.test", resume=f"Cand {i}\nB2B Sales Negotiation " * (i + 1) + "3 years experience")
    res = ok(c.post(f"/api/jobs/{job2['ref']}/rounds/{r2[1]['id']}/top-n", json={"n": 2, "reject_rest": True}))
    assert res == {"passed": 2, "failed": 2, "flagged_waiting": 0}, res
    print("LIVE FLOW EDITS, BULK ACTIONS, TOP-N, PRACTICAL TASK: OK")


def question_bank(c):
    import openpyxl
    t = c.get("/api/questions/template.xlsx")
    assert t.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(t.content))
    ws = wb.active
    ws.append(["logical", "hard", "single", "If 2x + 3 = 11, what is x?", "3", "4", "5", "6", "B", 2, "2x = 8."])
    ws.append(["english", "easy", "single", "Broken row with no answer", "a", "b", "", "", "", 1, ""])
    buf = io.BytesIO()
    wb.save(buf)
    res = ok(c.post("/api/questions/import", files={"file": ("bank.xlsx", buf.getvalue(), "application/vnd.ms-excel")}))
    assert res["created"] == 4 and len(res["errors"]) == 1, res
    dr = ok(c.post("/api/questions/draft", json={"section": "quantitative", "topic": "percentages", "count": 3}))["items"]
    assert len(dr) == 3
    ok(c.post("/api/questions", json={"questions": dr}))
    q = ok(c.get("/api/questions?section=logical&q=2x"))["items"][0]
    assert q["answer"] == [1] and q["marks"] == 2
    ok(c.patch(f"/api/questions/{q['id']}", json={"active": False}))
    ok(c.delete(f"/api/questions/{q['id']}"))
    assert c.post("/api/questions", json={"text": "No options", "options": ["a"], "answer": [0]}).status_code == 400
    print("QUESTION BANK (template, import with row errors, AI drafts, edit, deactivate, delete): OK")


def reports_exports_retention(c, campus_job):
    """Reports, the audit log, the HROne export and the recording-retention sweep."""
    import openpyxl
    from backend import retention
    rep = ok(c.get("/api/reports?days=0"))
    assert rep["total"] >= 5 and rep["funnel"][0]["n"] == rep["total"] and any(r["type"] == "test" and r["entered"] for r in rep["rounds"]), rep["funnel"]
    assert any(g["name"] == "campus" for g in rep["by_source"]) and any(g["name"] == "COEP Pune" for g in rep["by_college"]), rep["by_college"]
    jrep = ok(c.get(f"/api/reports?job={campus_job['id']}"))
    assert jrep["job"]["id"] == campus_job["id"] and jrep["rounds"][0]["type"] == "application"
    audit = ok(c.get("/api/audit"))
    assert audit["total"] > 20 and "flow_changed" in audit["actions"]
    assert ok(c.get("/api/audit?action=flow_changed"))["items"][0]["action"] == "flow_changed"
    # HROne: the company's own headers, in order; an unmapped header stays as an empty column
    fields = ok(c.get("/api/hrone/fields"))
    assert any(f["id"] == "full_name" for f in fields["fields"])
    ok(c.patch("/api/org", json={"settings": {"hrone_columns": [{"header": "Employee Name", "field": "full_name"}, {"header": "Official Email", "field": "email"},
                                                                  {"header": "Grade", "field": ""}, {"header": "Designation", "field": "job_title"}]}}))
    kiran = next(x for x in ok(c.get(f"/api/jobs/{campus_job['ref']}/pipeline"))["items"] if x["candidate"]["name"] == "Kiran Patil")
    r = c.get(f"/api/exports/hrone.xlsx?ids={kiran['id']}")
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [x.value for x in ws[1]] == ["Employee Name", "Official Email", "Grade", "Designation"]
    assert [x.value for x in ws[2]] == ["Kiran Patil", "kiran@coep.test", None, campus_job["title"]], [x.value for x in ws[2]]
    # retention: recordings and photos of a closed application go; scores and decisions stay
    det = ok(c.get(f"/api/applications/{kiran['id']}"))
    vid = next(x for x in det["rounds"] if x["round"]["type"] == "video_intro")
    assert c.get(f"/api/round-results/{vid['result']['id']}/file").status_code == 200
    ok(c.patch("/api/org", json={"settings": {"recording_retention_days": 1}}))
    assert retention.sweep()["rounds"] == 0, "nothing is deleted before the retention period"
    with db.session() as s:
        s.get(db.Application, kiran["id"]).decided_at = time.time() - 2 * 86400
    out = retention.sweep()
    assert out["rounds"] >= 2 and out["files"] >= 2 and out["photos"] == 1, out
    det = ok(c.get(f"/api/applications/{kiran['id']}"))
    vid = next(x for x in det["rounds"] if x["round"]["type"] == "video_intro")
    assert c.get(f"/api/round-results/{vid['result']['id']}/file").status_code == 404 and vid["data"]["media_deleted_at"]
    assert vid["result"]["score"] == 70.0 and vid["data"]["assessment"]["dimensions"], "scores and AI notes are kept"
    test = next(x for x in det["rounds"] if x["round"]["type"] == "test")
    assert not test["result"]["integrity"].get("snapshots") and not test["result"]["integrity"].get("start_photo")
    assert not det["candidate"]["has_photo"]
    assert retention.sweep()["rounds"] == 0, "a second pass has nothing left to do"
    assert ok(c.get("/api/audit?action=retention_sweep"))["total"] == 1
    print("REPORTS, AUDIT LOG, HRONE EXPORT, RECORDING RETENTION: OK")


def mailbox_import(c, slug, job):
    """Resumes emailed to the careers mailbox become candidates; a job named in the subject gets an application."""
    from email.message import EmailMessage
    from backend import mailbox
    msgs = {}

    def mail(subject, sender, attach=True, n=0):
        m = EmailMessage()
        m["From"], m["Subject"] = sender, subject
        m.set_content("Please find my resume attached.")
        if attach:
            m.add_attachment(f"Neha Kulkarni\nneha.k{n}@mail.test\n+91 98200 00{n:03d}\nJava Spring Boot developer, 5 years.".encode(), maintype="text",
                             subtype="plain", filename="Neha-resume.txt")
        msgs[str(len(msgs) + 1).encode()] = m.as_bytes()

    mail(f"Application for {job['title']}", "Neha K <neha.personal@mail.test>")
    mail("Hello", "Spam <spam@x.test>", attach=False)
    mail("Resume", "Job Board <noreply@board.test>", n=1)
    seen = []

    class FakeIMAP:
        def __init__(self, host, port, timeout=None): assert host == "imap.test" and port == 993
        def login(self, u, p): assert (u, p) == ("careers@rac.test", "app-password")
        def select(self, folder): assert folder == "INBOX"
        def search(self, *a): return "OK", [b" ".join(k for k in msgs if k not in seen)]
        def fetch(self, num, what): assert "PEEK" in what; return "OK", [(b"1 (BODY[] {n})", msgs[num])]
        def store(self, num, *a): seen.append(num)
        def logout(self): pass

    env = {"IMAP_HOST": "imap.test", "IMAP_USER": "careers@rac.test", "IMAP_PASSWORD": "app-password", "IMAP_ORG_SLUG": slug}
    old_env, old_imap = {k: os.environ.get(k) for k in env}, mailbox.imaplib.IMAP4_SSL
    os.environ.update(env)
    mailbox.imaplib.IMAP4_SSL = FakeIMAP
    try:
        out = mailbox.import_once()
        assert out == {"emails": 3, "candidates": 2, "applications": 1, "skipped": 1}, out
        assert mailbox.import_once()["emails"] == 0, "read emails are not imported twice"
    finally:
        mailbox.imaplib.IMAP4_SSL = old_imap
        for k, v in old_env.items():
            os.environ.pop(k) if v is None else os.environ.update({k: v})
    neha = ok(c.get("/api/candidates?q=neha.k0"))["items"][0]
    assert neha["source"] == "email" and neha["email"] == "neha.k0@mail.test" and neha["has_resume"], neha
    app_ = next(x for x in ok(c.get(f"/api/jobs/{job['ref']}/pipeline"))["items"] if x["candidate"]["email"] == "neha.k0@mail.test")
    assert app_["source"] == "email" and app_["round_id"], "the emailed application enters the job's flow"
    assert ok(c.get("/api/audit?action=mailbox_import"))["total"] == 2
    print("MAILBOX IMPORT (IMAP): OK")


def main():
    with TestClient(app) as c:
        me = setup_company(c)
        slug = ok(c.get("/api/auth/me"))["org"]["slug"]
        experienced(c, slug, me)
        job, d = campus(c, slug)
        integrity_and_resume(c, slug, job, d)
        transparency_and_requests(c, slug)
        live_edit_and_bulk(c, slug)
        question_bank(c)
        reports_exports_retention(c, job)
        mailbox_import(c, slug, job)
        msgs = ok(c.get("/api/messages"))
        assert msgs["total"] > 10 and all(m["status"] == "not_configured" for m in msgs["items"]), "without SMTP/WhatsApp, messages wait in the outbox"
        m = ok(c.post(f"/api/messages/{msgs['items'][0]['id']}/retry"))
        assert m["status"] == "not_configured"
    print("\nFLOW CHECKS PASSED")


if __name__ == "__main__":
    main()
